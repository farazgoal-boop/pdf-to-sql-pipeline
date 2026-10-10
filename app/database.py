"""Database models and helpers (SQLAlchemy 2.x). Works on SQLite, PostgreSQL and MySQL."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote

import pandas as pd
from sqlalchemy import (
    DateTime,
    Float,
    ForeignKey,
    Integer,
    String,
    Text,
    UniqueConstraint,
    create_engine,
    event,
    func,
    select,
)
from sqlalchemy.dialects import mysql, postgresql
from sqlalchemy.engine import Engine, make_url
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, relationship, sessionmaker
from sqlalchemy.schema import CreateIndex, CreateTable

from app.config import PROJECT_ROOT, settings

# TEXT is limited to 64 KB on MySQL; use LONGTEXT there (up to 4 GB).
LongText = Text().with_variant(mysql.LONGTEXT(), "mysql", "mariadb")

_MYSQL_TABLE_ARGS = {"mysql_engine": "InnoDB", "mysql_charset": "utf8mb4", "mysql_collate": "utf8mb4_unicode_ci"}


class Base(DeclarativeBase):
    pass


class Document(Base):
    __tablename__ = "documents"
    __table_args__ = (_MYSQL_TABLE_ARGS,)

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    file_name: Mapped[str] = mapped_column(String(255), nullable=False, index=True)
    file_path: Mapped[str] = mapped_column(String(1024), nullable=False)
    file_hash: Mapped[str] = mapped_column(String(64), nullable=False, unique=True)
    page_count: Mapped[int] = mapped_column(Integer, nullable=False, default=0)
    ocr_page_count: Mapped[int] = mapped_column(Integer, nullable=False, default=0)
    status: Mapped[str] = mapped_column(String(16), nullable=False)  # success / partial / failed
    error_message: Mapped[str | None] = mapped_column(Text, nullable=True)
    processed_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True), nullable=False, default=lambda: datetime.now(timezone.utc)
    )

    pages: Mapped[list[Page]] = relationship(
        back_populates="document",
        cascade="all, delete-orphan",
        passive_deletes=True,
        order_by="Page.page_number",
    )


class Page(Base):
    __tablename__ = "pages"
    __table_args__ = (
        # Also serves as the index for lookups by document_id.
        UniqueConstraint("document_id", "page_number", name="uq_pages_document_page"),
        _MYSQL_TABLE_ARGS,
    )

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    document_id: Mapped[int] = mapped_column(
        ForeignKey("documents.id", ondelete="CASCADE"), nullable=False
    )
    page_number: Mapped[int] = mapped_column(Integer, nullable=False)
    text: Mapped[str] = mapped_column(LongText, nullable=False, default="")
    extraction_method: Mapped[str] = mapped_column(String(8), nullable=False)  # text / ocr / error
    char_count: Mapped[int] = mapped_column(Integer, nullable=False, default=0)
    ocr_confidence: Mapped[float | None] = mapped_column(Float, nullable=True)

    document: Mapped[Document] = relationship(back_populates="pages")


def _enable_sqlite_foreign_keys(engine: Engine) -> None:
    @event.listens_for(engine, "connect")
    def _set_pragma(dbapi_connection, _record):  # pragma: no cover - trivial
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()


def mask_url(url: str) -> str:
    """Return the URL with any password replaced by ***."""
    try:
        # unquote only for readability (e.g. "C%3A" -> "C:"); the password is already masked.
        return unquote(make_url(url).render_as_string(hide_password=True))
    except Exception:
        return "<invalid DATABASE_URL>"


class Database:
    """Thin wrapper around a SQLAlchemy engine with the queries the app needs."""

    def __init__(self, url: str | None = None):
        self.url = url or settings.database_url
        parsed = make_url(self.url)
        self.backend = parsed.get_backend_name()  # sqlite / postgresql / mysql
        if self.backend == "sqlite" and parsed.database and parsed.database != ":memory:":
            db_file = Path(parsed.database)
            if not db_file.is_absolute():  # relative paths are relative to the project, not the CWD
                db_file = PROJECT_ROOT / db_file
                self.url = parsed.set(database=str(db_file)).render_as_string(hide_password=False)
            db_file.parent.mkdir(parents=True, exist_ok=True)

        self.engine = create_engine(self.url, pool_pre_ping=True)
        if self.backend == "sqlite":
            _enable_sqlite_foreign_keys(self.engine)
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(self.engine, expire_on_commit=False)

    # ------------------------------------------------------------------ info
    @property
    def masked_url(self) -> str:
        return mask_url(self.url)

    @property
    def location(self) -> str:
        parsed = make_url(self.url)
        if self.backend == "sqlite":
            return str(Path(parsed.database).resolve()) if parsed.database else ":memory:"
        port = f":{parsed.port}" if parsed.port else ""
        return f"{parsed.host}{port}/{parsed.database}"

    @property
    def display_location(self) -> str:
        """Short, safe-to-show location: SQLite path relative to the project, or the masked URL."""
        if self.backend != "sqlite":
            return self.masked_url
        if self.location == ":memory:":
            return self.location
        path = Path(self.location)
        try:
            return path.relative_to(PROJECT_ROOT.resolve()).as_posix()
        except ValueError:  # database file lives outside the project
            return path.as_posix()

    # ----------------------------------------------------------------- write
    def find_by_hash(self, file_hash: str) -> Document | None:
        with self.Session() as session:
            return session.scalar(select(Document).where(Document.file_hash == file_hash))

    def save_document(self, document: Document, replace: bool = False) -> Document:
        """Insert a document (with its pages) in a single transaction.

        If replace=True, an existing record with the same file_hash is deleted
        first, inside the same transaction, so a failure leaves the old record intact.
        """
        with self.Session.begin() as session:
            if replace:
                existing = session.scalar(select(Document).where(Document.file_hash == document.file_hash))
                if existing is not None:
                    session.delete(existing)
                    session.flush()
            session.add(document)
        return document

    # ------------------------------------------------------------------ read
    def stats(self) -> dict:
        with self.Session() as session:
            by_status = dict(
                session.execute(select(Document.status, func.count()).group_by(Document.status)).all()
            )
            page_row = session.execute(
                select(
                    func.count(Page.id),
                    func.coalesce(func.sum(Page.char_count), 0),
                    func.avg(Page.ocr_confidence),
                )
            ).one()
            ocr_pages = session.scalar(select(func.count()).where(Page.extraction_method == "ocr"))
        return {
            "documents": sum(by_status.values()),
            "success": by_status.get("success", 0),
            "partial": by_status.get("partial", 0),
            "failed": by_status.get("failed", 0),
            "pages": int(page_row[0]),
            "characters": int(page_row[1]),
            "ocr_pages": int(ocr_pages or 0),
            "avg_ocr_confidence": round(float(page_row[2]), 1) if page_row[2] is not None else None,
        }

    def documents_df(self) -> pd.DataFrame:
        query = select(
            Document.id,
            Document.file_name,
            Document.status,
            Document.page_count,
            Document.ocr_page_count,
            Document.processed_at,
            Document.error_message,
            Document.file_path,
            Document.file_hash,
        ).order_by(Document.id)
        with self.engine.connect() as conn:
            return pd.read_sql(query, conn)

    def pages_df(self, document_id: int | None = None) -> pd.DataFrame:
        query = (
            select(
                Document.file_name,
                Page.page_number,
                Page.extraction_method,
                Page.char_count,
                Page.ocr_confidence,
                Page.text,
                Page.document_id,
            )
            .join(Document, Page.document_id == Document.id)
            .order_by(Document.id, Page.page_number)
        )
        if document_id is not None:
            query = query.where(Page.document_id == document_id)
        with self.engine.connect() as conn:
            return pd.read_sql(query, conn)

    def get_page(self, document_id: int, page_number: int) -> Page | None:
        with self.Session() as session:
            return session.scalar(
                select(Page).where(Page.document_id == document_id, Page.page_number == page_number)
            )

    def search(self, term: str, limit: int = 100) -> list[dict]:
        """Case-insensitive substring search across all page text (portable LIKE)."""
        term = term.strip()
        if not term:
            return []
        query = (
            select(Document.id, Document.file_name, Page.page_number, Page.extraction_method, Page.text)
            .join(Document, Page.document_id == Document.id)
            .where(func.lower(Page.text).contains(term.lower(), autoescape=True))
            .order_by(Document.file_name, Page.page_number)
            .limit(limit)
        )
        with self.engine.connect() as conn:
            return [dict(row._mapping) for row in conn.execute(query)]

    # ---------------------------------------------------------------- export
    def export_pages_df(self) -> pd.DataFrame:
        """Every page with its file name and import time, ready for CSV/Excel."""
        query = (
            select(
                Document.file_name,
                Page.page_number,
                Page.extraction_method,
                Page.char_count,
                Page.ocr_confidence,
                Document.processed_at,
                Page.text,
            )
            .join(Document, Page.document_id == Document.id)
            .order_by(Document.id, Page.page_number)
        )
        with self.engine.connect() as conn:
            return _format_timestamps(pd.read_sql(query, conn))

    def export_documents_df(self) -> pd.DataFrame:
        return _format_timestamps(self.documents_df())

    def export_csv(self, path: str | Path) -> int:
        """Write every page to a CSV file that opens cleanly in Excel. Returns the row count."""
        df = self.export_pages_df()
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(to_csv_bytes(df))
        return len(df)

    def export_xlsx(self, path: str | Path) -> int:
        """Write a workbook with a Pages and a Documents sheet. Returns the page row count."""
        pages = self.export_pages_df()
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(to_xlsx_bytes({"Pages": pages, "Documents": self.export_documents_df()}))
        return len(pages)


# ---------------------------------------------------------------------- export helpers
EXCEL_CELL_LIMIT = 32_767  # maximum characters Excel can hold in one cell
_TRUNCATION_NOTE = " [...truncated for Excel; full text is in the database and CSV export]"
_WRAPPED_COLUMNS = {"text": 100, "error_message": 60}  # column -> width; wrapped and top-aligned


def _format_timestamps(df: pd.DataFrame) -> pd.DataFrame:
    """Render processed_at as a plain 'YYYY-MM-DD HH:MM:SS' (UTC) string so Excel shows it as-is."""
    if "processed_at" in df.columns:
        # Naive values (SQLite, MySQL) are stored in UTC; aware values (PostgreSQL) are converted to UTC.
        stamps = pd.to_datetime(df["processed_at"], utc=True)
        df["processed_at"] = stamps.dt.strftime("%Y-%m-%d %H:%M:%S").fillna("")
    return df


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    """UTF-8 with a byte-order mark, so Excel detects the encoding (accents, €, etc.)."""
    return df.to_csv(index=False).encode("utf-8-sig")


def _excel_safe(value):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(value) > EXCEL_CELL_LIMIT:
        value = value[: EXCEL_CELL_LIMIT - len(_TRUNCATION_NOTE)] + _TRUNCATION_NOTE
    return value


def to_xlsx_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Build an .xlsx: bold frozen header row, autofilter, sized columns, wrapped long-text columns."""
    from io import BytesIO

    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df = df.copy()
            for col in df.columns:
                if pd.api.types.is_string_dtype(df[col]) or df[col].dtype == object:
                    df[col] = df[col].map(_excel_safe)
            df.to_excel(writer, sheet_name=name, index=False)

            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            top = Alignment(vertical="top")
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = top
            for idx, col in enumerate(df.columns, start=1):
                letter = get_column_letter(idx)
                wrapped = col in _WRAPPED_COLUMNS
                if wrapped:
                    ws.column_dimensions[letter].width = _WRAPPED_COLUMNS[col]
                else:
                    longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(500) if pd.notna(v)])
                    ws.column_dimensions[letter].width = min(max(longest + 2, 10), 50)
                # Top-align every data cell so short values line up with the start of long wrapped text.
                data_alignment = Alignment(vertical="top", wrap_text=True) if wrapped else top
                for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    cell.alignment = data_alignment
    return buffer.getvalue()


def generate_schema_sql(path: str | Path = PROJECT_ROOT / "schema.sql") -> Path:
    """Write the CREATE TABLE statements for PostgreSQL and MySQL to schema.sql."""
    tables = [Document.__table__, Page.__table__]
    parts = [
        "-- PDF -> SQL Extractor: database schema",
        "-- Generated from app/database.py (SQLAlchemy models). The pipeline creates these",
        "-- tables automatically; this file is for reference or manual setup.",
        "",
    ]
    for title, dialect in (("PostgreSQL", postgresql.dialect()), ("MySQL / MariaDB (utf8mb4)", mysql.dialect())):
        parts.append(f"-- {'=' * 70}\n-- {title}\n-- {'=' * 70}\n")
        for table in tables:
            ddl = str(CreateTable(table).compile(dialect=dialect)).strip()
            parts.append("\n".join(line.rstrip() for line in ddl.splitlines()) + ";\n")
            for index in sorted(table.indexes, key=lambda i: i.name):
                parts.append(str(CreateIndex(index).compile(dialect=dialect)).strip() + ";\n")
    path = Path(path)
    path.write_text("\n".join(parts), encoding="utf-8")
    return path
